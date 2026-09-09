"""
End-to-end tests. Each one runs the real pipeline on generated data whose answers are known,
so a change to the matching logic that starts crediting the wrong family fails the build.

    pytest -q

The guarantee being protected is precision, not coverage: it is fine for the engine to send more
lines to a human, and never fine for it to allocate one to the wrong family.
"""
import csv, re, subprocess, sys, unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
AUTO_TIERS = {"M", "A", "B"}
SEEDS = [20260101, 7, 99]
# generate_fake_data.py:97 always bills one child 18.00 for supplies instead of 25.00.
PLANTED_SUPPLIES_ERROR = 7.00
# That planted error, plus any family invoiced at a sibling tier it no longer qualifies for
# because a child left. How many of those exist depends on the seed's enrolment, so this is a
# ceiling. A fee engine that is actually broken misses most of the roster, not a handful.
MAX_FEE_MISMATCHES = 5


def run(*args):
    r = subprocess.run([sys.executable, *map(str, args)], cwd=ROOT, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    return r.stdout


@pytest.fixture(scope="session", params=SEEDS)
def pipeline(request, tmp_path_factory):
    out = tmp_path_factory.mktemp(f"seed{request.param}")
    run(ROOT / "generate_fake_data.py", out, "--seed", request.param)
    ledger = out / "fee_ledger.xlsx"
    run(ROOT / "build_ledger.py", out / "roster.xlsx", out / "bank.xlsx", ledger)
    run(ROOT / "reconcile.py", ledger)

    wb = load_workbook(ledger, data_only=True)
    fam = {r[0]: str(r[1] or "") for r in wb["Families"].iter_rows(min_row=4, values_only=True) if r[0]}
    rows = []
    for r in wb["Receipts"].iter_rows(min_row=4, values_only=True):
        if r[0] is None:
            break
        rows.append(dict(
            date=str(r[0])[:10],
            amount=round(float(r[1] or 0), 2),
            payer=str(r[2] or "").strip().upper()[:23],
            tier=r[8] or "",
            fid=r[9] or "",
            cands=str(r[10] or ""),
            why=str(r[11] or ""),
        ))
    truth = defaultdict(list)
    for t in csv.DictReader(open(out / "ground_truth.csv")):
        key = (str(t["date"])[:10], round(float(t["amount"]), 2), t["payer"].strip().upper()[:23])
        truth[key].append(t)
    pairs = []
    for g in rows:
        bucket = truth[(g["date"], g["amount"], g["payer"])]
        assert bucket, f"receipt with no ground truth: {g}"
        pairs.append((g, bucket.pop(0)))
    return dict(wb=wb, fam=fam, pairs=pairs, seed=request.param)


def norm_family(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s.upper().replace(" / ", "-").replace("/", "-")).strip()


def correct(fam, fid, true_family):
    """Exact match on the family name, never substring.

    MARCHETTI is a substring of ACHTERBERG-MARCHETTI and they are unrelated families, so a
    substring test cannot see a receipt credited from one to the other - which is the precise
    error this whole suite exists to catch.
    """
    return bool(fid) and norm_family(fam.get(fid, "")) == norm_family(true_family)


def test_no_receipt_is_credited_to_the_wrong_family(pipeline):
    bad = [(g, t) for g, t in pipeline["pairs"]
           if g["tier"] in AUTO_TIERS and t["true_family"] and not correct(pipeline["fam"], g["fid"], t["true_family"])]
    assert not bad, f"{len(bad)} misallocated, first: {bad[0]}"


def test_most_receipts_are_allocated_without_a_human(pipeline):
    auto = [g for g, _ in pipeline["pairs"] if g["tier"] in AUTO_TIERS]
    assert len(auto) / len(pipeline["pairs"]) >= 0.75


def test_a_clean_reference_always_allocates(pipeline):
    for g, t in pipeline["pairs"]:
        if t["failure_mode"] == "clean":
            assert g["tier"] == "A", f"clean reference not matched by reference: {g}"


def test_a_reference_that_contradicts_the_payer_name_is_held_back(pipeline):
    for g, t in pipeline["pairs"]:
        if t["failure_mode"].startswith("wrong_ref") and g["tier"] in AUTO_TIERS:
            assert correct(pipeline["fam"], g["fid"], t["true_family"]), f"followed a wrong reference: {g}"


def test_conflicts_are_explained(pipeline):
    for g, _ in pipeline["pairs"]:
        if g["tier"] == "X":
            assert "may have typed the wrong invoice number" in g["why"]


def test_an_invoice_number_shared_by_two_families_never_guesses(pipeline):
    for g, t in pipeline["pairs"]:
        if t["failure_mode"].startswith("shared_invoice") and g["tier"] in AUTO_TIERS:
            assert correct(pipeline["fam"], g["fid"], t["true_family"]), f"guessed on a shared invoice number: {g}"


def test_every_queued_receipt_carries_a_reason(pipeline):
    for g, _ in pipeline["pairs"]:
        if g["tier"] not in AUTO_TIERS:
            assert g["why"].strip(), f"no reason given for a queued receipt: {g}"


def fee_inputs(wb):
    """The fee rules as the workbook states them, so this test checks the wiring and not a copy."""
    rates, terms = wb["Rates"], wb["Terms"]
    return dict(
        sibling_rate={n: rates.cell(row=3 + n, column=3).value for n in (1, 2, 3, 4)},
        wednesday_rate=rates.cell(row=9, column=3).value,
        reg_fee=rates.cell(row=11, column=3).value,
        supplies_fee=rates.cell(row=12, column=3).value,
        sessions={"Saturday": terms.cell(row=5, column=3).value,
                  "Wednesday": terms.cell(row=5, column=4).value},
    )


def expected_charge(row, siblings, f):
    """Recompute a child's expected fee in Python from the static columns.

    Columns I, K, L, M, R and T on the Children sheet are Excel formulas. openpyxl writes them
    without a cached value, so reading them with data_only=True yields None and any assertion
    over them is vacuous. Only the blue input columns are safe to read.
    """
    cls = row[4]
    n = row[9] if isinstance(row[9], (int, float)) else f["sessions"][cls]
    if siblings == 0:
        rate = 0.0
    elif cls == "Wednesday":
        rate = f["wednesday_rate"]
    else:
        rate = f["sibling_rate"][min(siblings, 4)] / siblings
    addons = str(row[14] or "")
    return round(
        round(rate * n, 2)
        - (row[13] or 0.0)
        + (f["reg_fee"] if "reg" in addons else 0.0)
        + (f["supplies_fee"] if "supplies" in addons else 0.0),
        2,
    )


def test_the_fee_rules_reproduce_the_invoices(pipeline):
    ch = pipeline["wb"]["Children"]
    enrolled = [r for r in ch.iter_rows(min_row=4, values_only=True) if r[0] and r[5] == "Enrolled"]
    assert enrolled
    f = fee_inputs(pipeline["wb"])
    assert f["sibling_rate"][1] and f["wednesday_rate"] and f["sessions"]["Saturday"], \
        "fee inputs missing from Rates/Terms - the rules are not readable as data"

    siblings = Counter(r[1] for r in enrolled)
    mismatched = []
    for r in enrolled:
        invoiced = r[18]
        assert isinstance(invoiced, (int, float)), f"{r[0]} is enrolled with no invoiced amount"
        want = expected_charge(r, siblings[r[1]], f)
        if abs(want - invoiced) >= 0.02:
            mismatched.append((r[0], want, invoiced))

    # The rules must reproduce essentially every invoice. A fee engine that has been broken -
    # a flat rate, a dropped sibling tier, the wrong session count - mismatches most of the roster,
    # so the ceiling is what gives this test teeth.
    assert len(mismatched) <= MAX_FEE_MISMATCHES, \
        f"fee rules disagree with {len(mismatched)} of {len(enrolled)} invoices: {mismatched[:5]}"

    # generate_fake_data.py:97 always bills one child 18.00 for supplies instead of 25.00. If that
    # stops being found, the check has quietly stopped checking - which is how this test used to
    # pass while asserting nothing.
    assert any(abs((want - got) - PLANTED_SUPPLIES_ERROR) < 0.01 for _, want, got in mismatched), \
        f"the planted {PLANTED_SUPPLIES_ERROR:.2f} supplies error was not detected: {mismatched}"


def test_no_money_disappears(pipeline):
    total = sum(g["amount"] for g, _ in pipeline["pairs"])
    allocated = sum(g["amount"] for g, _ in pipeline["pairs"] if g["tier"] in AUTO_TIERS)
    queued = sum(g["amount"] for g, _ in pipeline["pairs"] if g["tier"] not in AUTO_TIERS)
    assert round(allocated + queued, 2) == round(total, 2)


def test_a_different_term_needs_no_code_edit(tmp_path):
    """The school's next term must be a --term argument, not a change to six places in the source.

    build_ledger.py used to hard-code JAN-2026 and its date cut-offs, so running the tool for the
    term starting in September meant editing the file. This builds the same roster twice, for two
    different terms, and checks the workbook follows: the Terms sheet's current-term row, the
    session count the Children sheet looks up, and the term stamped on every child.
    """
    run(ROOT / "generate_fake_data.py", tmp_path, "--seed", 20260101)
    roster, bank = tmp_path / "roster.xlsx", tmp_path / "bank.xlsx"

    default_led, other_led = tmp_path / "default.xlsx", tmp_path / "other.xlsx"
    run(ROOT / "build_ledger.py", roster, bank, default_led)
    run(ROOT / "build_ledger.py", roster, bank, other_led, "--term", "AUT-2026")

    default_wb, other_wb = load_workbook(default_led), load_workbook(other_led)

    # Row 5 of Terms is the current term: the Children session formula and reconcile.py's ageing
    # formula both address it by row, so the two terms must land in the same place.
    assert default_wb["Terms"].cell(5, 1).value == "JAN-2026"
    assert other_wb["Terms"].cell(5, 1).value == "AUT-2026"
    # Row 4 is the prior term, which shifts along with it.
    assert default_wb["Terms"].cell(4, 1).value == "AUT-2025"
    assert other_wb["Terms"].cell(4, 1).value == "JAN-2026"
    # Sessions come from TERMS, not from a constant in the body of the script.
    assert default_wb["Terms"].cell(5, 3).value == 21
    assert other_wb["Terms"].cell(5, 3).value == 11
    # Every child is stamped with the term being built, and the Families header follows it.
    assert other_wb["Children"].cell(4, 7).value == "AUT-2026"
    assert "AUT-2026" in other_wb["Families"].cell(3, 5).value

    # An unknown term fails loudly rather than building a wrong workbook.
    bad = subprocess.run([sys.executable, str(ROOT / "build_ledger.py"), str(roster), str(bank),
                          str(tmp_path / "bad.xlsx"), "--term", "SPRING-1999"],
                         cwd=ROOT, capture_output=True, text=True)
    assert bad.returncode != 0 and "unknown term" in bad.stdout + bad.stderr


def test_the_workbook_has_the_sheets_a_reviewer_needs(pipeline):
    for sheet in ("Summary", "README", "Rates", "Terms", "Families", "Children", "Receipts", "Review", "Position"):
        assert sheet in pipeline["wb"].sheetnames
