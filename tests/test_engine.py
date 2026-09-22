"""
Unit tests for engine.py. No workbook, no subprocess: build a tiny roster in memory and check the
decisions directly. These run in milliseconds and pin the behaviours that matter most.

    pytest tests/test_engine.py -q
"""
from collections import Counter, defaultdict

import pytest

from engine import (Config, Receipt, Roster, allocate, decide, evidence, families_named,
                    norm, parse_refs, payer_key, surname_parts, tokens)

CFG = Config()


def roster(**fams):
    """roster(FAM_001=("HASANOVIC", "Elodie", "2026-001"), ...) -> Roster with one child per family."""
    r = Roster()
    firsts = Counter()
    for fid, (sur, first, inv) in fams.items():
        fid = fid.replace("_", "-")
        r.fam_name[fid] = sur.title()
        for part in surname_parts(sur):
            r.sur_index[part].add(fid)
        r.inv_index[inv].add(fid)
        r.expected[fid] = 409.5
        r.billed[fid] = 1
        firsts[first.upper()] += 1
        r.unique_first[first.upper()] = fid
    r.unique_first = {k: v for k, v in r.unique_first.items() if firsts[k] == 1}
    return r


def receipt(payer, memo, **kw):
    cur, prior = parse_refs(norm(memo), CFG)
    base = dict(r=4, date=None, amount=409.5, payer=payer, ref="", memo=memo, cur_ref=cur,
                prior_ref=prior, term=CFG.term_cur, override="", note="")
    base.update(kw)
    return Receipt(**base)


# ------------------------------------------------------------------ text helpers

def test_norm_strips_accents_and_uppercases():
    assert norm("Élodie Hasanović") == "ELODIE HASANOVIC"


def test_tokens_drop_stop_words_and_short_fragments():
    assert tokens("SCHOOL FEES for the CARDOSO family") == ["CARDOSO", "FAMILY"]


def test_payer_key_is_letters_only():
    assert payer_key("Mrs. O'Brien-Smith") == "MRSOBRIENSMITH"


def test_tokens_glue_a_leading_letter_but_never_a_possessive():
    """Only an apostrophe after a single leading letter is removed (M'BOLO, O'BRIEN), matching
    surname_parts(). A possessive or contraction stays a word boundary: ADAM'S must yield ADAM, or
    a memo for child Adam would manufacture an exact hit on the surname ADAMS. The typographic
    apostrophe counts too, since norm() would otherwise drop it and glue."""
    assert tokens("LOUIS M'BOLO O'BRIEN") == ["LOUIS", "MBOLO", "OBRIEN"]
    assert tokens("ADAM'S FEES") == ["ADAM"]
    assert tokens("ADAM\u2019S FEES") == ["ADAM"]
    assert tokens("MAALOUF'S FEES") == ["MAALOUF"]
    assert tokens("WE'LL PAY THE REST") == ["PAY", "REST"]


def test_apostrophe_surname_matches_exactly():
    """M'BOLO on the roster and M'BOLO in a memo must meet: both sides index as MBOLO."""
    assert surname_parts("M'BOLO") == ["MBOLO"]
    r = roster(FAM_001=("M'BOLO", "Louis", "2026-001"))
    assert r.sur_index["MBOLO"] == {"FAM-001"}
    tier, fid, cands, why = decide(evidence("LOUIS M'BOLO fees", r, CFG), r)
    assert (tier, fid) == ("B", "FAM-001")
    assert "surname MBOLO" in why


# ------------------------------------------------------------------ references

@pytest.mark.parametrize("memo, want", [
    ("2026-047", "2026-047"),
    ("2026047", "2026-047"),            # hyphen dropped
    ("MAALOU2026-030", "2026-030"),     # glued to the name
    ("26030", "2026-030"),              # two-digit year
    ("school fees", ""),
])
def test_parse_current_reference(memo, want):
    assert parse_refs(norm(memo), CFG)[0] == want


def test_prior_term_reference_only_when_no_current_one():
    assert parse_refs("2025-627", CFG) == ("", "2025-627")
    assert parse_refs("2025-627 2026-001", CFG) == ("2026-001", "")


def test_a_series_separates_two_terms_in_the_same_year():
    """Autumn 2026 issues 2026-5xx; the January 2026 term issued 2026-0xx. Same year, so the
    number range is what tells a current reference from a prior one."""
    cfg = Config(term_cur="AUT-2026", term_prior="JAN-2026", cur_series=(500, 699), prior_series=(1, 199))
    assert parse_refs("2026-612", cfg) == ("2026-612", "")
    assert parse_refs("HASANOVIC 2026-003", cfg) == ("", "2026-003")
    assert parse_refs("2026-003 2026-612", cfg) == ("2026-612", "")
    assert parse_refs("26612", cfg) == ("2026-612", "")            # two-digit year, still in series


def test_series_bounds_read_the_human_form():
    assert Config.series_bounds("2026-5xx / 2026-6xx") == (500, 699)
    assert Config.series_bounds("2026-0xx / 2026-1xx") == (0, 199)
    assert Config.series_bounds("2026-xxx") == () and Config.series_bounds(None) == ()


def test_no_series_means_the_year_decides_as_before():
    cfg = Config(term_cur="JAN-2026", term_prior="AUT-2025")
    assert cfg.in_series("999", "cur") and cfg.in_series("000", "prior")
    assert parse_refs("2026-999", cfg) == ("2026-999", "")


def test_reconcile_reads_the_term_off_the_ledger():
    """The Terms sheet says which term a ledger was built for; reconcile.py must not assume one."""
    from openpyxl import Workbook
    from reconcile import config_for
    wb = Workbook(); ws = wb.active; ws.title = "Terms"
    ws.cell(4, 1, "JAN-2026"); ws.cell(4, 3, 21); ws.cell(4, 5, "2026-0xx / 2026-1xx")
    ws.cell(5, 1, "AUT-2026"); ws.cell(5, 3, 11); ws.cell(5, 5, "2026-5xx / 2026-6xx")
    rates = wb.create_sheet("Rates")                        # per-session family rates, rows 4-6
    for row, rate in ((4, 19.5), (5, 34.5), (6, 46.0)): rates.cell(row, 3, rate)
    cfg = config_for(wb)
    assert (cfg.term_cur, cfg.term_prior) == ("AUT-2026", "JAN-2026")
    assert cfg.cur_series == (500, 699) and cfg.prior_series == (0, 199)
    assert dict(cfg.fee_bases)["AUT-2026"] == (214.5, 379.5, 506.0)


# ------------------------------------------------------------------ evidence and decisions

def test_exact_surname_is_tier_b():
    r = roster(FAM_001=("HASANOVIC", "Elodie", "2026-001"))
    tier, fid, cands, why = decide(evidence("ELODIE HASANOVIC 060", r, CFG), r)
    assert (tier, fid) == ("B", "FAM-001")
    assert "surname HASANOVIC" in why


def test_two_families_same_surname_is_a_tie_sent_to_review():
    r = roster(FAM_042=("PELLETIER", "Anna", "2026-042"), FAM_045=("PELLETIER", "Marc", "2026-045"))
    tier, fid, cands, why = decide(evidence("MARGIT PELLETIER 2025-627", r, CFG), r)
    assert tier == "D" and fid == "" and cands == {"FAM-042", "FAM-045"}
    assert "several families fit" in why


def test_truncated_surname_is_fuzzy_tier_c_not_allocated():
    r = roster(FAM_010=("VANTERPOOL", "Zora", "2026-010"))
    tier, fid, cands, why = decide(evidence("VANTER school fees", r, CFG), r)
    assert tier == "C" and fid == "" and cands == {"FAM-010"}


def test_possessive_memo_does_not_become_another_familys_surname():
    """MARTIN'S FEES is the Martin family's, not the Martins family's; ADAM'S FEES, where Adam is
    a child of one family and ADAMS the surname of another, stays a tie sent to review."""
    r = roster(FAM_010=("MARTIN", "Maya", "2026-010"), FAM_011=("MARTINS", "Zak", "2026-011"))
    tier, fid, cands, why = decide(evidence("MARTIN'S FEES", r, CFG), r)
    assert (tier, fid, why) == ("B", "FAM-010", "surname MARTIN")
    r = roster(FAM_001=("ADAMS", "Zoe", "2026-001"), FAM_002=("OTHER", "Adam", "2026-002"))
    tier, fid, cands, why = decide(evidence("ADAM'S FEES", r, CFG), r)
    assert tier not in CFG.auto_tiers and fid == ""
    assert cands == {"FAM-001", "FAM-002"}


def test_unique_child_first_name_is_evidence_but_not_enough_to_allocate():
    r = roster(FAM_010=("VANTERPOOL", "Halima", "2026-010"))
    ev = evidence("Halima", r, CFG)
    assert ev["FAM-010"] == ["child first name Halima"]
    assert families_named(ev) == set()
    assert decide(ev, r)[0] == "C"


# ------------------------------------------------------------------ allocation: the invariants

def test_reference_contradicted_by_memo_is_refused_tier_x():
    """The case that matters: a reference pointing at family A while the memo names family B."""
    r = roster(FAM_048=("RASMUSSEN", "Ida", "2026-048"), FAM_049=("OTHER", "Bo", "2026-039"))
    [x] = allocate([receipt("MARGIT RASMUSSEN", "MARGIT RASMUSSEN 2026-039")], r, CFG)
    assert x.tier == "X" and x.fid == ""
    assert x.cands == {"FAM-048", "FAM-049"}
    assert "may have typed the wrong invoice number" in x.why


def test_reference_agreeing_with_memo_is_tier_a():
    r = roster(FAM_047=("WENDELIN", "Margit", "2026-047"))
    [x] = allocate([receipt("MARGIT WENDELIN", "MARGIT WENDELIN 2026-047")], r, CFG)
    assert (x.tier, x.fid) == ("A", "FAM-047")


# ------------------------------------------------------------------ strictly richer surname evidence

def test_hits_is_the_per_token_view_that_evidence_folds():
    """evidence() is hits() folded per family: one matching loop, two views of it. part is the
    roster surname part the token matched - the token itself when the hit is exact, and '' for a
    child's first name, which is not surname evidence and so cannot make one family's evidence
    contain another's."""
    from engine import by_family, hits
    r = roster(FAM_002=("NAKASHIMA-PELLETIER", "Ida", "2026-003"), FAM_038=("NAKASHIMA", "Bo", "2026-008"))
    found = hits("IDA NAKASHIMA-PELLET", r, CFG)
    assert ("NAKASHIMA", "NAKASHIMA", "FAM-002", "surname NAKASHIMA") in found
    assert ("NAKASHIMA", "NAKASHIMA", "FAM-038", "surname NAKASHIMA") in found
    assert ("PELLET", "PELLETIER", "FAM-002", "'PELLET' starts PELLETIER") in found
    assert ("IDA", "", "FAM-002", "child first name Ida") in found
    assert by_family(found) == evidence("IDA NAKASHIMA-PELLET", r, CFG)


def test_truncated_compound_surname_contradicts_a_shared_reference():
    """The bank truncates the payer field at 23 characters, so NAKASHIMA-PELLETIER arrives as
    NAKASHIMA-PELLET. The shared reference narrows to the NAKASHIMA family, which the memo does
    name - but PELLET starts PELLETIER, so the compound family's surname evidence strictly
    contains the referenced family's: the memo describes that family and more."""
    r = roster(FAM_002=("NAKASHIMA-PELLETIER", "Ida", "2026-003"),
               FAM_038=("NAKASHIMA", "Bo", "2026-008"),
               FAM_007=("BAPTISTE", "Ann", "2026-008"))
    [x] = allocate([receipt("MARGIT NAKASHIMA-PELLET", "MARGIT NAKASHIMA-PELLET 2026-008")], r, CFG)
    assert x.tier == "X" and x.fid == ""
    assert {"FAM-002", "FAM-038"} <= x.cands
    assert "may have typed the wrong invoice number" in x.why


def test_truncated_compound_surname_contradicts_an_unshared_reference():
    """The same hole with no shared reference: the number resolves to one family on its own."""
    r = roster(FAM_002=("NAKASHIMA-PELLETIER", "Ida", "2026-003"),
               FAM_038=("NAKASHIMA", "Bo", "2026-038"))
    [x] = allocate([receipt("MARGIT NAKASHIMA-PELLET", "MARGIT NAKASHIMA-PELLET 2026-038")], r, CFG)
    assert x.tier == "X" and x.fid == ""
    assert x.cands == {"FAM-002", "FAM-038"}


def test_full_compound_surname_contradicts_a_reference_to_one_of_its_parts():
    """Untruncated, both parts are exact hits. The family bearing both contradicts a reference to
    either one; the family bearing only the other part does not, because its evidence is different
    words, not more of the same ones."""
    r = roster(FAM_002=("NAKASHIMA-PELLETIER", "Ida", "2026-003"),
               FAM_038=("NAKASHIMA", "Bo", "2026-038"),
               FAM_043=("PELLETIER", "Cy", "2026-043"))
    [x] = allocate([receipt("MARGIT NAKASHIMA-PELLETIER", "MARGIT NAKASHIMA-PELLETIER 2026-038")], r, CFG)
    assert x.tier == "X" and x.fid == ""
    assert "FAM-002" in x.cands
    assert "FAM-043" not in x.cands


def test_a_contradicted_reference_still_teaches_the_payer_alias():
    """Holding one row back must not re-route the rest of the pass: the payer key still stands for
    the family its reference named, so the next row decides on exactly the aliases it did before.
    Here row 5 names QUILLIAM outright and still follows the alias, as it always has."""
    r = roster(FAM_001=("ACHTERBERG-MARCHETTI", "Ida", "2026-001"), FAM_002=("ACHTERBERG", "Bo", "2026-002"),
               FAM_003=("MARCHETTI", "Cy", "2026-003"), FAM_004=("QUILLIAM", "Kai", "2026-004"))
    rows = allocate([
        receipt("ADRIEN ACHTERBERG-MARCHETTI", "ADRIEN ACHTERBERG-MARCHETTI 2026-002", r=4),
        receipt("ADRIEN ACHTERBERG-MARCHETTI", "QUILLIAM school fees", r=5),
    ], r, CFG)
    assert (rows[1].tier, rows[1].fid) == ("A", "FAM-002")


def test_a_payers_own_truncated_surname_does_not_contradict_their_reference():
    """SANDOV NILS is the Sandoval family, truncated by the bank, quoting their own number. No
    other family's evidence contains theirs, so the reference stands and teaches the alias that
    carries their next payment - even one that mentions another family's surname."""
    r = roster(FAM_001=("SANDOVAL", "Ida", "2026-001"), FAM_002=("NAKASHIMA", "Bo", "2026-002"))
    rows = allocate([
        receipt("SANDOV NILS", "SANDOV NILS 2026-001", r=4),
        receipt("SANDOV NILS", "SANDOV NILS NAKASHIMA SCHOOL", r=5),
    ], r, CFG)
    assert [(x.tier, x.fid) for x in rows] == [("A", "FAM-001"), ("A", "FAM-001")]


def test_a_truncated_or_misspelt_surname_still_corroborates_its_own_reference():
    """Fuzzy evidence for the referenced family is still corroboration, not a contradiction."""
    r = roster(FAM_001=("WOJCIECHOWSKI", "Ada", "2026-001"), FAM_002=("DZIEDZIC", "Bo", "2026-002"))
    [x] = allocate([receipt("WOJCIECHOWSK ORLA", "WOJCIECHOWSK ORLA 2026-001")], r, CFG)
    assert (x.tier, x.fid) == ("A", "FAM-001")
    [y] = allocate([receipt("DZIEDZIK ORLA", "DZIEDZIK ORLA 2026-002")], r, CFG)
    assert (y.tier, y.fid) == ("A", "FAM-002")


def test_a_truncated_surname_then_another_familys_name_follows_the_alias():
    r = roster(FAM_001=("SMITHSON", "Bo", "2026-001"), FAM_002=("MARTINEZ", "Ada", "2026-002"))
    rows = allocate([
        receipt("SMITHSO BO", "SMITHSO BO 2026-001", r=4),
        receipt("SMITHSO BO", "SMITHSO BO MARTINEZ", r=5),
    ], r, CFG)
    assert [(x.tier, x.fid) for x in rows] == [("A", "FAM-001"), ("A", "FAM-001")]


def test_two_surnames_neither_containing_the_other_is_not_a_richer_contradiction():
    """SMITH and DUPONT are different words: neither explains the other's evidence, so a reference
    to DUPONT in a memo carrying both is corroborated, not contradicted. Only the row whose memo
    names SMITH alone against a DUPONT reference is tier X, on the older exact-surname test."""
    r = roster(FAM_001=("DUPONT", "Ida", "2026-001"), FAM_002=("SMITH", "Bo", "2026-002"))
    rows = allocate([
        receipt("JOHN SMITH", "JOHN SMITH DUPONT 2026-001", r=4),
        receipt("JOHN SMITH", "JOHN SMITH 2026-001", r=5),
        receipt("JOHN SMITH", "JOHN SMITH payment", r=6),
    ], r, CFG)
    assert [(x.tier, x.fid) for x in rows] == [("A", "FAM-001"), ("X", ""), ("A", "FAM-001")]


def test_a_first_name_brushing_another_surname_does_not_contradict_a_reference():
    """NILS starts NILSSON and MARIE starts MARIETTE, but neither carries the DZIEDZIC hit the
    reference leans on, so neither family's evidence contains the referenced family's."""
    r = roster(FAM_057=("DZIEDZIC", "Ada", "2026-057"), FAM_060=("NILSSON", "Ann", "2026-060"))
    [x] = allocate([receipt("MR NILS DZIEDZIC", "MR NILS DZIEDZIC 2026-057")], r, CFG)
    assert (x.tier, x.fid) == ("A", "FAM-057")
    r = roster(FAM_057=("DZIEDZIC", "Ada", "2026-057"), FAM_018=("MARIETTE", "Ann", "2026-018"))
    [y] = allocate([receipt("MARIE-CLAIRE DZIEDZIC", "MARIE-CLAIRE DZIEDZIC 2026-057")], r, CFG)
    assert (y.tier, y.fid) == ("A", "FAM-057")


def test_a_stray_word_starting_a_surname_does_not_contradict_a_reference():
    """MARCH starts MARCHETTI. That family's evidence does not contain the Achterbergs', so a
    memo saying MARCH FEES cannot hold up their reference."""
    r = roster(FAM_002=("ACHTERBERG", "Bo", "2026-002"), FAM_031=("MARCHETTI", "Cy", "2026-031"))
    [x] = allocate([receipt("ADRIEN ACHTERBERG", "ADRIEN ACHTERBERG 2026-002 MARCH FEES")], r, CFG)
    assert (x.tier, x.fid) == ("A", "FAM-002")


def test_equal_surname_evidence_is_settled_by_the_reference_not_contradicted():
    """SANDOVAL hits the Sandovals and the Sandoval Zabalas identically: equal, not strictly
    richer, so the reference still decides between them and the payment allocates."""
    r = roster(FAM_040=("SANDOVAL", "Hugo", "2026-040"), FAM_004=("SANDOVAL ZABALA", "Kai", "2026-004"),
               FAM_050=("ZABALA", "Nils", "2026-050"))
    [x] = allocate([receipt("HUGO SANDOVAL", "HUGO SANDOVAL 2026-040")], r, CFG)
    assert (x.tier, x.fid) == ("A", "FAM-040")


def test_reference_from_a_payer_matching_nobody_is_not_allocated():
    """A reference alone is not enough: nothing in the memo or the payer ties it to the family,
    so it goes to review, and no alias is learned that would drag the next instalment along."""
    r = roster(FAM_039=("SMITH", "Ann", "2026-039"))
    rows = allocate([
        receipt("GRANDMA JONES", "GRANDMA JONES 2026-039", r=4),
        receipt("GRANDMA JONES", "2nd payment", r=5),
    ], r, CFG)
    x = rows[0]
    assert x.tier == "C" and x.fid == "" and x.cands == {"FAM-039"}
    assert "nothing else in the memo supports it" in x.why
    assert rows[1].tier != "A" and rows[1].fid == ""


def test_wrong_reference_from_an_unknown_payer_is_held_back():
    """The misallocation this rule exists for: an unknown payer quoting someone else's invoice
    number. Tier X cannot see it (the memo names nobody), so corroboration must."""
    r = roster(FAM_048=("RASMUSSEN", "Ida", "2026-048"), FAM_049=("OTHER", "Bo", "2026-039"))
    [x] = allocate([receipt("GRANDMA JONES", "GRANDMA JONES 2026-039")], r, CFG)
    assert x.tier not in CFG.auto_tiers and x.fid == ""
    assert x.cands == {"FAM-049"}


def test_reference_corroborated_by_a_misspelt_surname_is_tier_a():
    """Fuzzy evidence cannot allocate on its own, but it does corroborate a reference."""
    r = roster(FAM_030=("MAALOUF", "Genevieve", "2026-030"))
    [x] = allocate([receipt("GENEVEVE MALOUF", "GENEVEVE MALOUF 2026-030")], r, CFG)
    assert (x.tier, x.fid) == ("A", "FAM-030")


def test_reference_corroborated_by_a_seeded_alias_is_tier_a():
    """A payer typed into the Families sheet for this family corroborates the reference."""
    r = roster(FAM_030=("MAALOUF", "Genevieve", "2026-030"))
    r.alias_seed["GRANDMAJONES"].add("FAM-030")
    [x] = allocate([receipt("GRANDMA JONES", "GRANDMA JONES 2026-030")], r, CFG)
    assert (x.tier, x.fid) == ("A", "FAM-030")


@pytest.mark.parametrize("payer", ["ABC", "A B", ""])
def test_short_payer_key_does_not_corroborate_a_reference(payer):
    """A key below min_alias_len is too ambiguous to be an alias, so it cannot vouch for a later
    bare reference either - otherwise every anonymous line would corroborate the next one."""
    r = roster(FAM_001=("HASANOVIC", "Elodie", "2026-001"))
    rows = allocate([
        receipt(payer, "HASANOVIC 2026-001", r=4),
        receipt(payer, "2026-001", r=5),
    ], r, CFG)
    assert (rows[0].tier, rows[0].fid) == ("A", "FAM-001")
    assert rows[1].tier == "C" and rows[1].fid == "" and rows[1].cands == {"FAM-001"}


def test_payer_known_for_two_families_does_not_corroborate_a_reference():
    """A grandparent paying for two families: their key points at both, so it vouches for neither
    a bare reference nor a bare payment."""
    r = roster(FAM_001=("HASANOVIC", "Elodie", "2026-001"), FAM_002=("KOWALSKI", "Jan", "2026-002"))
    rows = allocate([
        receipt("GRANDMA JONES", "HASANOVIC 2026-001", r=4),
        receipt("GRANDMA JONES", "KOWALSKI 2026-002", r=5),
        receipt("GRANDMA JONES", "2026-002", r=6),
        receipt("GRANDMA JONES", "final instalment", r=7),
    ], r, CFG)
    assert [(x.tier, x.fid) for x in rows[:2]] == [("A", "FAM-001"), ("A", "FAM-002")]
    assert rows[2].tier == "C" and rows[2].fid == "" and rows[2].cands == {"FAM-002"}
    assert rows[3].tier == "D" and rows[3].fid == ""


def test_held_back_reference_to_another_family_blocks_the_alias():
    """A payer confirmed for one family who then quotes another family's invoice number: the
    conflicting line is held for review, and so is their next bare payment - the alias must not
    quietly pick the first family while the engine has just flagged a second one."""
    r = roster(FAM_040=("SMITH", "Ann", "2026-040"), FAM_041=("OTHER", "Bo", "2026-041"))
    rows = allocate([
        receipt("NANA P", "SMITH 2026-040", r=4),
        receipt("NANA P", "2026-041", r=5),
        receipt("NANA P", "final instalment", r=6),
    ], r, CFG)
    assert (rows[0].tier, rows[0].fid) == ("A", "FAM-040")
    assert rows[1].tier == "C" and rows[1].fid == "" and rows[1].cands == {"FAM-041"}
    assert rows[2].tier == "D" and rows[2].fid == ""


def test_unsupported_reference_still_shows_the_family_the_memo_does_fit():
    """Fuzzy evidence for a different family is not a contradiction (that is tier X, exact surname
    only), but the reviewer should see it next to the held-back reference."""
    r = roster(FAM_010=("VANTERPOOL", "Halima", "2026-010"), FAM_011=("QUILLIAM", "Ulrich", "2026-011"))
    [x] = allocate([receipt("SOMEONE", "ULRICH 2026-010")], r, CFG)
    assert x.tier == "C" and x.fid == "" and x.cands == {"FAM-010"}
    assert "nothing else in the memo supports it" in x.why
    assert x.why.endswith("  -  memo also fits FAM-011 (child first name Ulrich)")


def test_alias_learned_from_reference_matches_a_later_bare_payment():
    """Confirm a payer once by reference; their next payment with no reference at all follows."""
    r = roster(FAM_030=("MAALOUF", "Genevieve", "2026-030"))
    rows = allocate([
        receipt("GENEVIEVE MAALOUF", "MAALOU2026-030", r=4),
        receipt("GENEVIEVE MAALOUF", "2nd payment", r=5),
    ], r, CFG)
    assert [(x.tier, x.fid) for x in rows] == [("A", "FAM-030"), ("A", "FAM-030")]
    assert "previously seen with a verified reference" in rows[1].why


def test_alias_is_not_learned_from_an_ambiguous_payer_key():
    """A very short payer key must not become an alias, whatever it paid for."""
    r = roster(FAM_001=("HASANOVIC", "Elodie", "2026-001"))
    rows = allocate([
        receipt("ABC", "HASANOVIC 2026-001", r=4),
        receipt("ABC", "no reference here", r=5),
    ], r, CFG)
    assert rows[0].tier == "A"
    assert rows[1].tier == "D"          # nothing usable, and no alias to fall back on


def test_manual_override_wins_and_teaches_an_alias():
    r = roster(FAM_001=("HASANOVIC", "Elodie", "2026-001"))
    rows = allocate([
        receipt("GRANDMA JONES", "gift", override="FAM-001", r=4),
        receipt("GRANDMA JONES", "another gift", r=5),
    ], r, CFG)
    assert (rows[0].tier, rows[0].fid) == ("M", "FAM-001")
    assert (rows[1].tier, rows[1].fid) == ("A", "FAM-001")


def test_shared_invoice_number_resolved_by_surname_in_memo():
    r = roster(FAM_008=("CARDOSO", "Ulrich", "2026-008"), FAM_045=("PELLETIER", "Marc", "2026-008"))
    [x] = allocate([receipt("ULRICH CARDOSO", "CARDOSO 2026-008")], r, CFG)
    assert (x.tier, x.fid) == ("A", "FAM-008")


def test_shared_invoice_number_with_no_surname_is_review():
    r = roster(FAM_008=("CARDOSO", "Ulrich", "2026-008"), FAM_045=("PELLETIER", "Marc", "2026-008"))
    [x] = allocate([receipt("SOMEONE ELSE", "2026-008")], r, CFG)
    assert x.tier == "D" and x.cands == {"FAM-008", "FAM-045"}


def test_precision_invariant_nothing_auto_allocated_without_exact_evidence():
    """Property: every automatically allocated receipt has an override, a reference, an alias,
    or an exact surname behind it. Fuzzy evidence alone never allocates."""
    r = roster(FAM_010=("VANTERPOOL", "Zora", "2026-010"), FAM_011=("QUILLIAM", "Ulrich", "2026-011"))
    rows = allocate([
        receipt("A B", "VANTERPOL fees", r=4),        # misspelt
        receipt("C D", "Zora", r=5),                   # child's name only
        receipt("E F", "QUILLIAM school fees", r=6),   # exact surname
    ], r, CFG)
    auto = [x for x in rows if x.tier in CFG.auto_tiers]
    assert [x.fid for x in auto] == ["FAM-011"]
    for x in rows:
        if x.tier in CFG.auto_tiers:
            assert x.override or x.cur_ref or "surname " in x.why or "previously seen" in x.why
